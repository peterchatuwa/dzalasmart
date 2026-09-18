#!/usr/bin/env python3
"""
Generate Word document with Airtel USSD Peer IP information
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_page_break(doc):
    """Add a page break"""
    doc.add_page_break()

def add_horizontal_line(paragraph):
    """Add a horizontal line to a paragraph"""
    p = paragraph._element
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '000000')
    pBdr.append(bottom)
    pPr.append(pBdr)

def create_peer_ip_document():
    """Create the Word document"""
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    # ============================================================================
    # TITLE PAGE
    # ============================================================================
    
    # Add some space
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Main title
    title = doc.add_heading('AIRTEL USSD INTEGRATION', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.runs[0]
    title_run.font.color.rgb = RGBColor(0, 0, 128)
    title_run.font.size = Pt(26)
    title_run.bold = True
    
    # Subtitle
    subtitle = doc.add_heading('Peer IP Address & Technical Specifications', level=2)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.runs[0]
    subtitle_run.font.color.rgb = RGBColor(0, 0, 128)
    subtitle_run.font.size = Pt(18)
    
    doc.add_paragraph()
    
    # Service name
    service = doc.add_paragraph('Nzeru za Alimi')
    service.alignment = WD_ALIGN_PARAGRAPH.CENTER
    service_run = service.runs[0]
    service_run.font.size = Pt(16)
    service_run.italic = True
    
    service2 = doc.add_paragraph('(Smart Farming Platform)')
    service2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    service2_run = service2.runs[0]
    service2_run.font.size = Pt(14)
    service2_run.italic = True
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Date
    date_para = doc.add_paragraph('Document Date: September 17, 2026')
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_para.runs[0]
    date_run.font.size = Pt(12)
    
    # Prepared for
    prepared = doc.add_paragraph('Prepared for: Airtel Malawi USSD Integration Team')
    prepared.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prepared_run = prepared.runs[0]
    prepared_run.font.size = Pt(12)
    
    add_page_break(doc)
    
    # ============================================================================
    # SECTION 1: PEER IP ADDRESS (MAIN INFO)
    # ============================================================================
    
    heading = doc.add_heading('1. VPN PEER IP ADDRESS', 1)
    heading_run = heading.runs[0]
    heading_run.font.color.rgb = RGBColor(0, 0, 128)
    
    doc.add_paragraph(
        'This is the IP address that Airtel\'s USSD gateway should connect to.'
    )
    
    # Create a table for the main IP info
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Light Grid Accent 1'
    
    # Header row
    header_cells = table.rows[0].cells
    header_cells[0].text = 'Parameter'
    header_cells[1].text = 'Value'
    
    # Make header bold
    for cell in header_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(12)
    
    # Data rows
    data = [
        ('Public IP Address', '37.60.252.211'),
        ('Domain Name', 'api.zammunda.com'),
        ('Port', '443 (HTTPS)'),
        ('Protocol', 'TCP')
    ]
    
    for i, (param, value) in enumerate(data, start=1):
        row = table.rows[i]
        row.cells[0].text = param
        row.cells[1].text = value
        
        # Bold the value column
        for paragraph in row.cells[1].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0, 100, 0)
    
    doc.add_paragraph()
    
    # Important note
    note = doc.add_paragraph()
    note.add_run('IMPORTANT: ').bold = True
    note.add_run('This is a direct HTTPS API connection, not a VPN tunnel. ')
    note.add_run('No VPN concentrator or IPSec configuration is required.')
    
    doc.add_paragraph()
    
    # ============================================================================
    # SECTION 2: API ENDPOINT DETAILS
    # ============================================================================
    
    heading = doc.add_heading('2. API ENDPOINT DETAILS', 1)
    heading_run = heading.runs[0]
    heading_run.font.color.rgb = RGBColor(0, 0, 128)
    
    # USSD Endpoint
    doc.add_heading('2.1 USSD Endpoint', 2)
    
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Light Grid Accent 1'
    
    header_cells = table.rows[0].cells
    header_cells[0].text = 'Parameter'
    header_cells[1].text = 'Value'
    for cell in header_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    endpoint_data = [
        ('Full URL', 'https://api.zammunda.com/ussd'),
        ('HTTP Method', 'POST'),
        ('Content-Type', 'application/x-www-form-urlencoded')
    ]
    
    for i, (param, value) in enumerate(endpoint_data, start=1):
        row = table.rows[i]
        row.cells[0].text = param
        row.cells[1].text = value
    
    doc.add_paragraph()
    
    # Request Parameters
    doc.add_heading('2.2 Request Parameters', 2)
    
    params = [
        ('sessionId', 'Session identifier (provided by Airtel)', 'String', 'Required'),
        ('serviceCode', 'USSD short code (e.g., *413#)', 'String', 'Required'),
        ('phoneNumber', 'User\'s phone number (E.164 format)', 'String', 'Required'),
        ('text', 'User\'s input text', 'String', 'Optional')
    ]
    
    table = doc.add_table(rows=len(params) + 1, cols=4)
    table.style = 'Light Grid Accent 1'
    
    header = table.rows[0].cells
    header[0].text = 'Parameter'
    header[1].text = 'Description'
    header[2].text = 'Type'
    header[3].text = 'Required'
    
    for cell in header:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    for i, (name, desc, ptype, req) in enumerate(params, start=1):
        row = table.rows[i]
        row.cells[0].text = name
        row.cells[1].text = desc
        row.cells[2].text = ptype
        row.cells[3].text = req
    
    doc.add_paragraph()
    
    # Response Format
    doc.add_heading('2.3 Response Format', 2)
    
    table = doc.add_table(rows=3, cols=2)
    table.style = 'Light Grid Accent 1'
    
    header = table.rows[0].cells
    header[0].text = 'Response Type'
    header[1].text = 'Description'
    for cell in header:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    responses = [
        ('CON <message>', 'Continue session - Display menu or prompt to user'),
        ('END <message>', 'End session - Display final message and terminate')
    ]
    
    for i, (rtype, desc) in enumerate(responses, start=1):
        row = table.rows[i]
        row.cells[0].text = rtype
        row.cells[1].text = desc
    
    doc.add_paragraph()
    
    # ============================================================================
    # SECTION 3: DNS CONFIGURATION
    # ============================================================================
    
    heading = doc.add_heading('3. DNS CONFIGURATION', 1)
    heading_run = heading.runs[0]
    heading_run.font.color.rgb = RGBColor(0, 0, 128)
    
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Light Grid Accent 1'
    
    header = table.rows[0].cells
    header[0].text = 'DNS Parameter'
    header[1].text = 'Value'
    for cell in header:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    dns_data = [
        ('Domain Name', 'api.zammunda.com'),
        ('DNS A Record', '37.60.252.211'),
        ('TTL (Time To Live)', '300 seconds (5 minutes)')
    ]
    
    for i, (param, value) in enumerate(dns_data, start=1):
        row = table.rows[i]
        row.cells[0].text = param
        row.cells[1].text = value
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.add_run('Verification Command: ').bold = True
    code_p1 = doc.add_paragraph('nslookup api.zammunda.com')
    code_p1.paragraph_format.left_indent = Inches(0.5)
    for run in code_p1.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(10)
    code_p2 = doc.add_paragraph('dig api.zammunda.com')
    code_p2.paragraph_format.left_indent = Inches(0.5)
    for run in code_p2.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(10)
    
    doc.add_paragraph()
    
    # ============================================================================
    # SECTION 4: SECURITY & SSL/TLS
    # ============================================================================
    
    heading = doc.add_heading('4. SECURITY & SSL/TLS CONFIGURATION', 1)
    heading_run = heading.runs[0]
    heading_run.font.color.rgb = RGBColor(0, 0, 128)
    
    table = doc.add_table(rows=7, cols=2)
    table.style = 'Light Grid Accent 1'
    
    header = table.rows[0].cells
    header[0].text = 'Security Parameter'
    header[1].text = 'Status / Value'
    for cell in header:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    security_data = [
        ('SSL Certificate Status', '✓ Valid & Active'),
        ('Certificate Type', 'Let\'s Encrypt (Auto-renewed every 90 days)'),
        ('TLS Version', 'TLS 1.2 / TLS 1.3'),
        ('Certificate Expiry', 'Auto-renewed (No manual intervention required)'),
        ('IP Whitelisting', '✓ Active (Only Airtel gateway IPs accepted)'),
        ('Rate Limiting', '✓ Active (10 requests/minute per phone number)')
    ]
    
    for i, (param, value) in enumerate(security_data, start=1):
        row = table.rows[i]
        row.cells[0].text = param
        row.cells[1].text = value
    
    doc.add_paragraph()
    
    # ============================================================================
    # SECTION 5: NETWORK TOPOLOGY
    # ============================================================================
    
    heading = doc.add_heading('5. NETWORK TOPOLOGY', 1)
    heading_run = heading.runs[0]
    heading_run.font.color.rgb = RGBColor(0, 0, 128)
    
    doc.add_paragraph('Connection Flow Diagram:')
    
    flow = [
        '1. User dials *413# on their mobile phone',
        '2. Request reaches Airtel Mobile Network',
        '3. Airtel USSD Gateway (SOURCE) processes the request',
        '   Internal Gateway IPs: 172.26.166.100-103, 172.26.166.41-43, etc.',
        '4. Gateway sends HTTPS POST request over the internet',
        '   TO: https://api.zammunda.com/ussd (37.60.252.211)',
        '5. Your VPS Server (DESTINATION/PEER) receives and processes',
        '6. Server responds with CON or END message',
        '7. Response travels back through Airtel Network to user\'s phone'
    ]
    
    for step in flow:
        p = doc.add_paragraph(step, style='List Bullet')
        p.paragraph_format.left_indent = Inches(0.5)
    
    doc.add_paragraph()
    
    note = doc.add_paragraph()
    note.add_run('Key Points:').bold = True
    doc.add_paragraph('• SOURCE = Airtel USSD Gateway (initiates connection)', style='List Bullet')
    doc.add_paragraph('• DESTINATION/PEER = Your VPS Server (receives connection)', style='List Bullet')
    doc.add_paragraph('• Connection Type = Standard HTTPS over internet (Port 443)', style='List Bullet')
    doc.add_paragraph('• No VPN tunnel required', style='List Bullet')
    
    doc.add_paragraph()
    
    # ============================================================================
    # SECTION 6: CONNECTIVITY TESTING
    # ============================================================================
    
    heading = doc.add_heading('6. CONNECTIVITY TESTING', 1)
    heading_run = heading.runs[0]
    heading_run.font.color.rgb = RGBColor(0, 0, 128)
    
    doc.add_paragraph(
        'Airtel can perform the following tests to verify connectivity before going live:'
    )
    
    doc.add_heading('6.1 Health Check Endpoint', 2)
    p = doc.add_paragraph()
    p.add_run('URL: ').bold = True
    p.add_run('https://api.zammunda.com/health')
    
    p = doc.add_paragraph()
    p.add_run('Command: ').bold = True
    code_p = doc.add_paragraph('curl https://api.zammunda.com/health')
    code_p.paragraph_format.left_indent = Inches(0.5)
    for run in code_p.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(10)
    
    p = doc.add_paragraph()
    p.add_run('Expected Response: ').bold = True
    resp_p = doc.add_paragraph('{"ok":true,"service":"nzeru-za-alimi","uptime":...}')
    resp_p.paragraph_format.left_indent = Inches(0.5)
    for run in resp_p.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(10)
    
    doc.add_paragraph()
    
    doc.add_heading('6.2 USSD Endpoint Test', 2)
    p = doc.add_paragraph()
    p.add_run('Command: ').bold = True
    
    test_cmd = '''curl -X POST https://api.zammunda.com/ussd \\
  -d "sessionId=test123&serviceCode=*413#&phoneNumber=%2B265888000001&text="'''
    code_p = doc.add_paragraph(test_cmd)
    code_p.paragraph_format.left_indent = Inches(0.5)
    for run in code_p.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
    
    p = doc.add_paragraph()
    p.add_run('Expected Response: ').bold = True
    resp_p = doc.add_paragraph('CON Enter your PIN:')
    resp_p.paragraph_format.left_indent = Inches(0.5)
    for run in resp_p.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(10)
    
    doc.add_paragraph()
    
    doc.add_heading('6.3 SSL Certificate Verification', 2)
    p = doc.add_paragraph()
    p.add_run('Command: ').bold = True
    code_p = doc.add_paragraph(
        'openssl s_client -connect api.zammunda.com:443 -servername api.zammunda.com'
    )
    code_p.paragraph_format.left_indent = Inches(0.5)
    for run in code_p.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
    
    doc.add_paragraph()
    
    # ============================================================================
    # SECTION 7: FIREWALL REQUIREMENTS
    # ============================================================================
    
    heading = doc.add_heading('7. FIREWALL REQUIREMENTS (AIRTEL SIDE)', 1)
    heading_run = heading.runs[0]
    heading_run.font.color.rgb = RGBColor(0, 0, 128)
    
    doc.add_paragraph(
        'Airtel should configure their firewall to allow outbound connections:'
    )
    
    table = doc.add_table(rows=6, cols=2)
    table.style = 'Light Grid Accent 1'
    
    header = table.rows[0].cells
    header[0].text = 'Firewall Parameter'
    header[1].text = 'Value'
    for cell in header:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    firewall_data = [
        ('Direction', 'OUTBOUND (from Airtel to Your Server)'),
        ('Source IP', 'Airtel USSD Gateway (172.26.x.x - internal)'),
        ('Destination IP', '37.60.252.211 (Your VPS Server)'),
        ('Protocol', 'TCP'),
        ('Port', '443 (HTTPS)')
    ]
    
    for i, (param, value) in enumerate(firewall_data, start=1):
        row = table.rows[i]
        row.cells[0].text = param
        row.cells[1].text = value
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.add_run('Action: ').bold = True
    p.add_run('ALLOW')
    
    doc.add_paragraph()
    
    # ============================================================================
    # SECTION 8: CONTACT INFORMATION
    # ============================================================================
    
    heading = doc.add_heading('8. CONTACT INFORMATION', 1)
    heading_run = heading.runs[0]
    heading_run.font.color.rgb = RGBColor(0, 0, 128)
    
    doc.add_heading('8.1 Technical Contact', 2)
    
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Light Grid Accent 1'
    
    header = table.rows[0].cells
    header[0].text = 'Contact Field'
    header[1].text = 'Information'
    for cell in header:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    contact_data = [
        ('Name', 'Peter Chatuwa'),
        ('Organization', 'Nzeru za Alimi / Zammunda'),
        ('Role', 'System Administrator / Technical Lead'),
        ('Email', 'peter.chatuwa@zammunda.com')
    ]
    
    for i, (field, info) in enumerate(contact_data, start=1):
        row = table.rows[i]
        row.cells[0].text = field
        row.cells[1].text = info
    
    doc.add_paragraph()
    
    doc.add_heading('8.2 Server Hosting', 2)
    
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Light Grid Accent 1'
    
    header = table.rows[0].cells
    header[0].text = 'Hosting Parameter'
    header[1].text = 'Details'
    for cell in header:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    hosting_data = [
        ('Server Location', 'Malawi / Southern Africa'),
        ('Public IP Address', '37.60.252.211'),
        ('Uptime SLA', '99.5%')
    ]
    
    for i, (param, detail) in enumerate(hosting_data, start=1):
        row = table.rows[i]
        row.cells[0].text = param
        row.cells[1].text = detail
    
    doc.add_paragraph()
    
    # ============================================================================
    # SECTION 9: QUICK REFERENCE SUMMARY
    # ============================================================================
    
    heading = doc.add_heading('9. QUICK REFERENCE SUMMARY', 1)
    heading_run = heading.runs[0]
    heading_run.font.color.rgb = RGBColor(128, 0, 0)
    
    doc.add_paragraph(
        'When Airtel asks for "Peer IP" or "VPN Peer IP Address", provide:'
    )
    
    # Create highlighted box
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Medium Shading 1 Accent 1'
    
    quick_ref = [
        ('IP Address', '37.60.252.211'),
        ('Domain', 'api.zammunda.com'),
        ('Port', '443'),
        ('Protocol', 'HTTPS (TCP)')
    ]
    
    for i, (key, value) in enumerate(quick_ref):
        row = table.rows[i]
        row.cells[0].text = key
        row.cells[1].text = value
        
        # Make value bold and larger
        for paragraph in row.cells[1].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(14)
                run.font.color.rgb = RGBColor(0, 100, 0)
    
    doc.add_paragraph()
    
    # Final note
    note_box = doc.add_paragraph()
    note_box.add_run('IMPORTANT CLARIFICATION:').bold = True
    doc.add_paragraph(
        'When asked "Is this a VPN?", the answer is: NO, this is a direct HTTPS API '
        'connection. No VPN tunnel, no IPSec, no VPN concentrator required. '
        'This is standard internet HTTPS traffic over port 443.'
    )
    
    doc.add_paragraph()
    
    # ============================================================================
    # SECTION 10: AVAILABILITY & SLA
    # ============================================================================
    
    heading = doc.add_heading('10. AVAILABILITY & SERVICE LEVEL', 1)
    heading_run = heading.runs[0]
    heading_run.font.color.rgb = RGBColor(0, 0, 128)
    
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Light Grid Accent 1'
    
    header = table.rows[0].cells
    header[0].text = 'Service Parameter'
    header[1].text = 'Commitment'
    for cell in header:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    sla_data = [
        ('Server Availability', '24/7 (99.5% uptime SLA)'),
        ('Response Time', '< 2 seconds (95th percentile)'),
        ('Session Timeout', '180 seconds (3 minutes)'),
        ('Maintenance Window', 'Sundays 2:00-4:00 AM CAT (if needed)')
    ]
    
    for i, (param, value) in enumerate(sla_data, start=1):
        row = table.rows[i]
        row.cells[0].text = param
        row.cells[1].text = value
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # ============================================================================
    # FOOTER
    # ============================================================================
    
    footer_para = doc.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.add_run(
        '─────────────────────────────────────────────────\n'
        'End of Document\n'
        'Prepared for Airtel Malawi USSD Integration Team\n'
        'September 17, 2026'
    )
    footer_run.font.size = Pt(10)
    footer_run.font.color.rgb = RGBColor(128, 128, 128)
    footer_run.italic = True
    
    # Save document
    output_path = '/workspace/AIRTEL_USSD_PEER_IP_INFORMATION.docx'
    doc.save(output_path)
    print(f'✓ Word document created: {output_path}')
    print(f'✓ File size: {os.path.getsize(output_path):,} bytes')
    return output_path

if __name__ == '__main__':
    import os
    create_peer_ip_document()
